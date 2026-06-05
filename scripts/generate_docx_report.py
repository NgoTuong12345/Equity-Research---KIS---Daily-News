#!/usr/bin/env python3
import sys
import json
import glob
from pathlib import Path
import docx
import unicodedata

sys.path.insert(0, str(Path(__file__).resolve().parent))
from report_paths import BASE_DIR, export_file, find_data_file, find_existing, REPORTS_DIR
from news_title_rules import normalized_item_title, ticker_company_exchange_title

def strip_accents(text: str) -> str:
    return "".join(c for c in unicodedata.normalize('NFD', text) if unicodedata.category(c) != 'Mn').replace('đ', 'd').replace('Đ', 'D')

def generate_report(base: str, lang: str):
    TEMPLATE_DIR = BASE_DIR / "news_html_template"
    
    parts = base.split("_")
    filename_date = f"{parts[1]}.{parts[2]}.{parts[3]}"
    filename_session = "Afternoon News" if "after" in base else "Morning News"
    
    if lang == "en":
        out_filename = f"{filename_session}_KIS RESEARCH_{filename_date}.docx"
    else:
        out_filename = f"{filename_session}_KIS RESEARCH_{filename_date}_vn.docx"
        
    out_path = export_file(base, "docx", out_filename)
    print(f"Generating {lang.upper()} DOCX report: {out_filename}")
    
    # Load JSON files
    trading_file = find_data_file(base, "trading")
    macro_file = find_data_file(base, "macro")
    corp_file = find_data_file(base, "corporate")
    epo_file = find_data_file(base, "economy_political_others")
    
    trading_data = json.loads(trading_file.read_text(encoding="utf-8")) if trading_file.exists() else {"items": []}
    macro_data = json.loads(macro_file.read_text(encoding="utf-8")) if macro_file.exists() else {"items": []}
    corp_data = json.loads(corp_file.read_text(encoding="utf-8")) if corp_file.exists() else {"sectors": []}
    epo_data = json.loads(epo_file.read_text(encoding="utf-8")) if epo_file.exists() else {"subtypes": []}
    
    # Open the existing docx to use as a template (to preserve margins and styling defaults)
    fallback_path = TEMPLATE_DIR / "Afternoon News_KIS RESEARCH_04.06.2026.docx"
    if fallback_path.exists():
        doc = docx.Document(str(fallback_path))
    else:
        doc = docx.Document()
            
    # Clear all paragraphs in the document
    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)
        
    # Translate and format helper for Date
    months_en = ["January", "February", "March", "April", "May", "June",
                 "July", "August", "September", "October", "November", "December"]
    try:
        day, month, year = int(parts[1]), int(parts[2]), int(parts[3])
        if lang == "en":
            date_str = f"{months_en[month - 1]} {day:02d}, {year}"
        else:
            date_str = f"ngày {day:02d} tháng {month:02d} năm {year}"
    except Exception:
        date_str = "June 05, 2026" if lang == "en" else "ngày 05 tháng 06 năm 2026"
        
    if lang == "en":
        session_label = "afternoon" if "after" in base else "morning"
        dear_str = "Dear Sir, "
        intro_prefix = "Please kindly have the update by "
        intro_suffix = " as listed below:"
        h1_trading = "1. Trading news:"
        h1_macro = "2. Macro indicators:"
        h1_corp = "3. Vietnamese industry/ corporate news:"
        h1_epo = "4. Political, Social, and Economic News"
    else:
        session_label = "buổi chiều" if "after" in base else "buổi sáng"
        dear_str = "Kính gửi Quý khách, "
        intro_prefix = "Xin vui lòng xem thông tin cập nhật vào "
        intro_suffix = " dưới đây:"
        h1_trading = "1. Tin giao dịch:"
        h1_macro = "2. Các chỉ số vĩ mô:"
        h1_corp = "3. Tin tức ngành / doanh nghiệp Việt Nam:"
        h1_epo = "4. Tin tức Chính trị, Xã hội và Kinh tế"
        
    # Add paragraphs
    # 1. Dear Sir,
    doc.add_paragraph(dear_str)
    
    # 2. Intro paragraph
    p = doc.add_paragraph()
    p.add_run(intro_prefix)
    p.add_run(session_label).bold = True
    p.add_run(" (")
    p.add_run(date_str).bold = True
    p.add_run(")" + intro_suffix)
    
    # 3. 1. Trading news: (Bold)
    doc.add_paragraph().add_run(h1_trading).bold = True
    
    # Trading items
    for item in trading_data.get("items", []):
        title = ticker_company_exchange_title(item, lang)
        summary = item.get(f"summary_{lang}", "")
        doc.add_paragraph(f"{title}: {summary}" if title else summary)
        
    # 4. 2. Macro indicators: (Bold)
    doc.add_paragraph()
    doc.add_paragraph().add_run(h1_macro).bold = True
    
    # Macro items use the three approved title buckets instead of repeating "Macro".
    for item in macro_data.get("items", []):
        title = normalized_item_title(item, "macro", lang)
        summary = item.get(f"summary_{lang}", "")
        doc.add_paragraph(f"{title}: {summary}" if title else summary)
            
    # 5. 3. Vietnamese industry/ corporate news: (Bold)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run(h1_corp).bold = True
    
    SECTOR_ORDER = ["banking", "financials", "consumers", "industrials", "materials", "real_estate", "technologies", "pharma", "utilities", "oil_gas"]
    
    if lang == "en":
        SECTOR_MAP = {
            "banking": "[Banking]", "financials": "[Financials]", "consumers": "[Consumers]",
            "industrials": "[Industrials]", "materials": "[Basic Materials]", "real_estate": "[Real Estates]",
            "technologies": "[Technologies]", "pharma": "[Pharma]", "utilities": "[Utilities]", "oil_gas": "[Oil & Gas]"
        }
    else:
        SECTOR_MAP = {
            "banking": "[Ngân hàng]", "financials": "[Tài chính]", "consumers": "[Tiêu dùng]",
            "industrials": "[Công nghiệp]", "materials": "[Nguyên vật liệu]", "real_estate": "[Bất động sản]",
            "technologies": "[Công nghệ]", "pharma": "[Dược phẩm]", "utilities": "[Tiện ích]", "oil_gas": "[Dầu khí]"
        }
    
    sectors_by_key = {sec["sector_key"]: sec for sec in corp_data.get("sectors", [])}
    for sec_key in SECTOR_ORDER:
        sec = sectors_by_key.get(sec_key)
        if sec and sec.get("items"):
            doc.add_paragraph().add_run(SECTOR_MAP.get(sec_key, f"[{sec['section_title_' + lang]}]")).bold = True
            for item in sec["items"]:
                ticker = item.get("ticker", "")
                title = ticker_company_exchange_title(item, lang)
                summary = item.get(f"summary_{lang}", "")
                
                p_text = f"{title}: {summary}" if title else summary
                doc.add_paragraph(p_text)
                
    # 6. 4. Political, Social, and Economic News (Bold)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run(h1_epo).bold = True
    
    SUBTYPE_ORDER = ["commodities", "economies_investments", "political", "policies", "social", "international_relation", "others"]
    
    if lang == "en":
        SUBTYPE_MAP = {
            "commodities": "[Commodities Price]", "economies_investments": "[Economy & Investments]",
            "political": "[Political]", "policies": "[Policies]", "social": "[Social]",
            "international_relation": "[International Relations]", "others": "[Others]"
        }
    else:
        SUBTYPE_MAP = {
            "commodities": "[Giá hàng hóa]", "economies_investments": "[Kinh tế & Đầu tư]",
            "political": "[Chính trị]", "policies": "[Chính sách]", "social": "[Xã hội]",
            "international_relation": "[Quan hệ quốc tế]", "others": "[Khác]"
        }
    
    subtypes_by_key = {sub["subtype_key"]: sub for sub in epo_data.get("subtypes", [])}
    for sub_key in SUBTYPE_ORDER:
        sub = subtypes_by_key.get(sub_key)
        if sub and sub.get("items"):
            for item in sub["items"]:
                prefix = SUBTYPE_MAP.get(sub_key, f"[{sub['section_title_' + lang]}]")
                title = normalized_item_title(item, "economy_political_others", lang)
                summary = item.get(f"summary_{lang}", "")
                p_text = f"{prefix} {title}: {summary}" if title else f"{prefix} {summary}"
                doc.add_paragraph(p_text)
                
    # Save the generated document
    doc.save(str(out_path))
    print(f"Successfully generated DOCX at: {out_path.name}")

def main():
    base = "after_04_06_2026"
    if len(sys.argv) > 1:
        base = sys.argv[1]
    
    generate_report(base, "en")
    generate_report(base, "vn")

if __name__ == "__main__":
    main()
